"""
Analytics API endpoints — Exposes business intelligence data for Dashboard & AI Agent.
"""
import io
import csv
import logging
from typing import Any, Dict, Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.api.auth import get_current_user
from app.services.analytics_service import AnalyticsService

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/analytics", tags=["analytics"], dependencies=[Depends(get_current_user)])


@router.get("/overview", summary="Executive KPI summary & period-over-period comparison")
def get_overview(
    date_range: str = Query(default="30d", description="Preset: today, 7d, 30d, 3m, 6m, 12m, all, custom"),
    source_name: Optional[str] = Query(default=None, description="Filter by data source platform"),
    start_date: Optional[str] = Query(default=None, description="ISO start date if date_range=custom"),
    end_date: Optional[str] = Query(default=None, description="ISO end date if date_range=custom"),
    db: Session = Depends(get_db),
) -> Dict[str, Any]:
    svc = AnalyticsService(db)
    return svc.get_overview_kpis(date_range, source_name, start_date, end_date)


@router.get("/revenue-trends", summary="Daily/Periodic time-series revenue & order volume")
def get_revenue_trends(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    svc = AnalyticsService(db)
    return svc.get_revenue_trends(date_range, source_name, start_date, end_date)


@router.get("/sales-breakdown", summary="Sales breakdown by product category, platform, and top products")
def get_sales_breakdown(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    svc = AnalyticsService(db)
    return svc.get_sales_breakdown(date_range, source_name, start_date, end_date)


@router.get("/customers", summary="Customer growth, retention cohorts, LTV, and churn indicators")
def get_customer_analytics(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    svc = AnalyticsService(db)
    return svc.get_customer_analytics(date_range, source_name, start_date, end_date)


@router.get("/alerts", summary="Automated business anomaly signals & performance indicators")
def get_business_alerts(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    svc = AnalyticsService(db)
    return svc.get_business_alerts(date_range, source_name, start_date, end_date)


@router.get("/inventory-risk", summary="Stock risk based on imported on-hand quantities and recent sales")
def get_inventory_risk(
    source_name: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    return AnalyticsService(db).get_inventory_risk(source_name)


@router.get("/ai-context", summary="Structured analytics payload formatted for Person 3 AI Agent")
def get_ai_context(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    """
    Returns a unified JSON data package designed for Person 3's AI Agent
    to consume for natural language reasoning, anomaly analysis, and recommendation generation.
    """
    svc = AnalyticsService(db)
    return svc.get_ai_context(date_range, source_name, start_date, end_date)


@router.get("/export/csv", summary="Download aggregated business analytics report as CSV")
def export_csv_report(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    svc = AnalyticsService(db)
    trends = svc.get_revenue_trends(date_range, source_name, start_date, end_date)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Date", "Revenue (USD)", "Completed Orders", "Average Order Value (USD)"])

    for r in trends:
        writer.writerow([r["date"], r["revenue"], r["orders"], r["avg_order_value"]])

    output.seek(0)
    filename = f"analytics_report_{date_range}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/xlsx", summary="Download executive analytics workbook")
def export_xlsx_report(
    date_range: str = Query(default="30d"),
    source_name: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    svc = AnalyticsService(db)
    overview = svc.get_overview_kpis(date_range, source_name, start_date, end_date)
    trends = svc.get_revenue_trends(date_range, source_name, start_date, end_date)
    sales = svc.get_sales_breakdown(date_range, source_name, start_date, end_date)
    customers = svc.get_customer_analytics(date_range, source_name, start_date, end_date)
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Executive Summary"
    summary.append(["Clearview BI — Executive Summary"])
    summary.append(["Period", overview["period"]["preset"]])
    summary.append(["Currency", "USD (converted where a supported rate is available)"])
    summary.append(["Exchange-rate source", "ExchangeRate-API (daily reference rates)"])
    summary.append(["Metric", "Current", "Previous", "Change (%)"])
    for label, key in (("Revenue (USD)", "revenue"), ("Orders", "orders"), ("Average order value (USD)", "avg_order_value"), ("Active customers", "active_customers")):
        metric = overview["kpis"][key]
        summary.append([label, metric["current"], metric.get("previous"), metric.get("percentage_change")])
    summary.append(["Unconverted orders", overview.get("unconverted_orders_excluded", 0)])

    trend_sheet = workbook.create_sheet("Sales Trends")
    trend_sheet.append(["Date", "Revenue (USD)", "Orders", "Average order value (USD)"])
    for row in trends:
        trend_sheet.append([row.get("date"), row.get("revenue"), row.get("orders"), row.get("avg_order_value")])

    products_sheet = workbook.create_sheet("Top Products")
    products_sheet.append(["Product", "SKU", "Category", "Units sold", "Orders", "Revenue (USD)"])
    for row in sales.get("top_products", []):
        products_sheet.append([row.get("name"), row.get("sku"), row.get("category"), row.get("units_sold"), row.get("order_count"), row.get("revenue")])

    customer_sheet = workbook.create_sheet("Customer Segments")
    customer_sheet.append(["Segment", "Customers", "Share (%)", "Suggested action"])
    for segment in (customers.get("rfm_segmentation") or {}).get("segments", []):
        customer_sheet.append([segment.get("name"), segment.get("count"), segment.get("percentage"), segment.get("strategy")])

    try:
        from app.services.ai_agent_service import AIAgentService
        recommendations = AIAgentService(db).generate_recommendations(date_range=date_range, source_name=source_name)
        recommendation_sheet = workbook.create_sheet("AI Recommendations")
        recommendation_sheet.append(["Priority", "Recommendation", "Expected impact", "Action steps"])
        for recommendation in recommendations.get("recommendations", []):
            steps = "\n".join(recommendation.get("action_steps", []))
            recommendation_sheet.append([recommendation.get("priority"), recommendation.get("title"), recommendation.get("expected_impact"), steps])
    except Exception:
        logger.exception("Unable to include AI recommendations in workbook")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A5" if sheet.title == "Executive Summary" else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=13)
            cell.fill = PatternFill("solid", fgColor="246B4D")
        if sheet.title == "Executive Summary":
            for cell in sheet[5]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="397F60")
        for column in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 60)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=clearview_bi_report_{date_range}.xlsx"})
